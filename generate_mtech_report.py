import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_complete_document():
    doc = docx.Document()
    
    # Page setup: Left 1.5 in, others 1.0 in
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.5)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = False

    def set_font(run, size=12, bold=False, italic=False, color=None, font_name="Times New Roman"):
        run.font.name = font_name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5, bold=False, italic=False, size=12):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            set_font(r, size=size, bold=bold, italic=italic)
        return p

    def add_chapter_title(chap_num, chap_name):
        doc.add_page_break()
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(24)
        p1.paragraph_format.space_after = Pt(12)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(f"CHAPTER {chap_num}\n\n{chap_name.upper()}")
        set_font(r1, size=14, bold=True)
        return p1

    def add_sec_heading(sec_title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(sec_title)
        set_font(r, size=12, bold=True)
        return p

    def add_subsec_heading(subsec_title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(subsec_title)
        set_font(r, size=12, bold=True, italic=True)
        return p

    def add_body(text, indent=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        r = p.add_run(text)
        set_font(r, size=12)
        return p

    def add_bullet(text, level=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        r_bullet = p.add_run("•  ")
        set_font(r_bullet, size=12, bold=True)
        r = p.add_run(text)
        set_font(r, size=12)
        return p

    def add_fig(image_path, fig_num, caption):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            r_img = p_img.add_run()
            r_img.add_picture(image_path, width=Inches(5.4))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(f"Fig. {fig_num} {caption}")
            set_font(r_cap, size=10, bold=True)
        else:
            print(f"Warning: Figure {image_path} not found.")

    def format_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Header
        hdr_cells = table.rows[0].cells
        for i, h_text in enumerate(headers):
            hdr_cells[i].text = h_text
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_font(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        
        # Rows
        for r_idx, row_data in enumerate(data):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
                row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_font(run, size=9.5)
        
        # Set widths
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    def add_custom_table(tab_num, title, headers, data, col_widths):
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(12)
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run(f"Table {tab_num} {title}")
        set_font(r_cap, size=10, bold=True)
        
        t = doc.add_table(rows=len(data) + 1, cols=len(headers))
        format_table(t, col_widths, headers, data)
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(8)

    print("Building Document Sections...")
    return doc, set_font, add_p, add_chapter_title, add_sec_heading, add_subsec_heading, add_body, add_bullet, add_fig, add_custom_table

print("generate_mtech_report.py helper defined.")
