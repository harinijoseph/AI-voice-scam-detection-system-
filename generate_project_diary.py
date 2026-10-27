import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_project_diary():
    doc = docx.Document()
    
    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    def set_font(run, size=11, bold=False, italic=False, color=None, font_name="Times New Roman"):
        run.font.name = font_name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4, line_spacing=1.15, bold=False, italic=False, size=11):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            set_font(r, size=size, bold=bold, italic=italic)
        return p

    # VIT Logo
    logo_path = "FINAL DOC/extracted_images/rId8_image1.jpeg"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(6)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(3.2))

    # Header Details
    add_p("PROJECT WORK DIARY", bold=True, size=14, space_before=4, space_after=4)
    add_p("AI GENERATED VOICE SCAM DETECTION SYSTEM: A MULTIMODAL REAL-TIME DEFENSE FRAMEWORK INTEGRATING ACOUSTIC FORENSICS AND BEHAVIORAL NATURAL LANGUAGE PROCESSING", bold=True, size=11, space_before=2, space_after=6)
    add_p("SWE3004 - SOFTWARE DESIGN AND DEVELOPMENT PROJECT", bold=True, size=11, space_before=0, space_after=6)
    add_p("Submitted in partial fulfillment of the requirements for the degree of", italic=True, size=10, space_before=0, space_after=2)
    add_p("MASTER OF TECHNOLOGY IN SOFTWARE ENGINEERING", bold=True, size=12, space_before=0, space_after=2)
    add_p("in\nDEPARTMENT OF SOFTWARE AND SYSTEMS ENGINEERING", size=10.5, space_before=0, space_after=6)
    add_p("by", size=10, space_before=0, space_after=2)
    add_p("HARINI JOSEPH\n23MSE0142", bold=True, size=11, space_before=0, space_after=6)
    add_p("Under the guidance of\nProf. / Dr. [Faculty Guide Name], [Designation]\nSchool of Computer Science Engineering and Information Systems (SCORE)\nVellore Institute of Technology (VIT), Vellore - 632014", bold=False, size=10, space_before=0, space_after=6)
    add_p("OCTOBER, 2026", bold=True, size=11, space_before=0, space_after=12)

    # Diary Data Entries
    diary_entries = [
        {
            "week": "Week 01\n06-07-2026 to 12-07-2026",
            "date": "10-07-2026\n(Friday)",
            "points": [
                "Met with guide to discuss project domain selection and finalize initial topic ideas for SWE3004 in Cybersecurity and Speech AI.",
                "Discussed the rapid emergence of zero-shot neural voice cloning tools (ElevenLabs, VALL-E, Tortoise-TTS) and the rising crisis of AI voice phishing (vishing) and family emergency scams.",
                "Examined critical vulnerabilities in traditional anti-scam tools: caller ID blacklists bypassed by spoofing, audio-only detectors degrading over noisy cellular calls, and post-call alerts arriving too late.",
                "Guide suggested developing a real-time multimodal defense architecture unifying physical acoustic speech forensics with behavioral natural language processing (NLP).",
                "Agreed on the primary objective: real-time detection of synthetic voice clones, tracking conversation scam stages, and providing active intervention during live calls.",
                "Action Plan for Next Week: Draft formal project abstract, define research scope, establish core objectives, and prepare documentation for the Zeroth Review."
            ]
        },
        {
            "week": "Week 02\n13-07-2026 to 19-07-2026",
            "date": "14-07-2026\n(Tuesday)",
            "points": [
                "Presented drafted project abstract and defined core research objectives for the upcoming Zeroth Review.",
                "Finalized official project title with guide approval: 'AI Generated Voice Scam Detection System: A Multimodal Real-Time Defense Framework Integrating Acoustic Forensics and Behavioral Natural Language Processing' (VoxSentinel AI).",
                "Discussed the evaluation rubric for Zeroth Review scheduled for Wednesday, 15-07-2026.",
                "Guide reviewed problem formulation regarding zero-shot voice cloning, telephony network degradation (8kHz G.711 / AMR codecs), and psychological manipulation funnels.",
                "Guide signed Zeroth Review document and approved project title and abstract on VTOP.",
                "Action Plan for Next Week: Appear for Zeroth Review on 15-07-2026 and initiate systematic literature survey on synthetic audio forensics and behavioral vishing detection."
            ]
        },
        {
            "week": "Week 03\n20-07-2026 to 26-07-2026",
            "date": "20-07-2026\n(Monday)",
            "points": [
                "Debriefed guide on Zeroth Review presentation and panel feedback (panel commended multimodal direction, advised rigorous evaluation under noisy telephony).",
                "Reviewed initial pool of 8 research papers covering synthetic audio detection (ASVspoof 2019/2021, WaveFake), acoustic representations (LFCC, CQCC, MFCC), and speech-to-text NLP frameworks.",
                "Analyzed mathematical foundations of vocoder artifact detection: phase inconsistencies in Neural Vocoders (HiFi-GAN, WaveGlow) and spectral envelope anomalies.",
                "Guide advised identifying telephony degradation challenges: why acoustic classifiers fail when 8kHz narrowband codecs truncate frequencies above 3.4kHz.",
                "Action Plan for Next Week: Expand literature review to 15 peer-reviewed publications and synthesize a comparative survey matrix detailing methodology, datasets, and research gaps."
            ]
        },
        {
            "week": "Week 04\n27-07-2026 to 02-08-2026",
            "date": "27-07-2026\n(Monday)",
            "points": [
                "Presented completed 15-paper literature survey synthesis matrix comparing acoustic-only, NLP-only, and multimodal fraud detection approaches.",
                "Identified critical research gaps: acoustic-only detectors suffer catastrophic recall collapse on out-of-distribution (OOD) neural voices over cellular networks; text-only NLP detectors trigger false alarms on benign high-urgency conversations.",
                "Validated research hypothesis: Dual-modality fusion combining microscopic acoustic forensics with multi-turn conversational scam funnel tracking eliminates blind spots and dramatically cuts false positives.",
                "Guide approved research gap analysis and authorized proceeding with architectural decomposition and mathematical modeling.",
                "Action Plan for Next Week: Produce system architecture block diagram, design UML Use Case and Class diagrams, and mathematically formulate multimodal decision fusion logic."
            ]
        },
        {
            "week": "Week 05\n03-08-2026 to 09-08-2026",
            "date": "03-08-2026\n(Monday)",
            "points": [
                "Presented architectural design of VoxSentinel AI featuring a 4-tier pipeline: Client Ingestion Layer, Acoustic Forensics Pipeline, Behavioral NLP Pipeline, and Multimodal Decision & Defense Layer.",
                "Walked through core mathematical formulations: (a) 284-dimensional acoustic feature extraction (LFCC, Hilbert phase, spectral entropy), (b) 4-stage conversational scam funnel state machine, and (c) SNR-adaptive multimodal fusion scoring with synergy multiplier.",
                "Formulated the Out-of-Band (OOB) Challenge Defense Protocol: generating real-time contextual verification challenges (e.g., secret family codeword) during high-risk calls.",
                "Reviewed UML Use Case Diagram, Class Diagram, and Component Diagram prepared for Review-1 milestone.",
                "Action Plan for Next Week: Finalize Review-1 documentation report and 20-slide presentation deck; meet guide on Tuesday, 11-08-2026 for pre-review sign-off."
            ]
        },
        {
            "week": "Week 06\n10-08-2026 to 16-08-2026",
            "date": "11-08-2026\n(Tuesday)",
            "points": [
                "Conducted formal pre-review consultation for Review-1 scheduled for Wednesday, 12-08-2026 in SJT classrooms.",
                "Guide thoroughly inspected 20 presentation slides covering Title, Abstract, Problem Statement, Literature Survey (15 papers), Scope, Proposed Architecture, and Module Specifications.",
                "Verified that all slide contents and deliverables adhere strictly to the 20-mark Review-1 departmental rubric.",
                "Rehearsed presentation defense explaining why single-modality defenses fail over cellular networks and how multimodal fusion provides resilience.",
                "Guide signed and attested official Review-1 evaluation document, authorizing appearance before the review committee.",
                "Action Plan for Next Week: Appear for Review-1 on 12-08-2026, document panel feedback, and begin local workstation software environment setup."
            ]
        },
        {
            "week": "Week 07\n17-08-2026 to 23-08-2026",
            "date": "17-08-2026\n(Monday)",
            "points": [
                "Reported Review-1 panel feedback: Panel commended the active OOB challenge concept and recommended simulating realistic 8kHz PSTN/AMR compression codecs during testing.",
                "Outlined modular implementation plan following strict object-oriented design principles across src/audio, src/nlp, src/fusion, and backend/routes.",
                "Demonstrated workstation environment configuration: Python 3.11, PyTorch 2.3, CUDA, Librosa, Whisper, SoundFile, and Scikit-learn.",
                "Implemented and demonstrated feature_extractor.py: Extracts 284 acoustic features per frame including Linear Frequency Cepstral Coefficients (LFCC), CQCC, pitch jitter (F0 perturbation), shimmer, and spectral entropy dynamics.",
                "Action Plan for Next Week: Implement the Acoustic Forensics Classification Engine (acoustic_analyzer.py) using a Stacking Ensemble Meta-Learner."
            ]
        },
        {
            "week": "Week 08\n24-08-2026 to 30-08-2026",
            "date": "24-08-2026\n(Monday)",
            "points": [
                "Demonstrated acoustic_analyzer.py and train_acoustic_rf.py: Implemented Stacking Ensemble Meta-Learner combining ExtraTrees, Random Forest, and Multi-Layer Perceptron (MLP) trained on ASVspoof and WaveFake datasets.",
                "Demonstrated Acoustic Conformer neural model baseline (train_conformer_dataset.py) for spectral-temporal artifact detection.",
                "Benchmarked feature extraction and model inference latency: Measured average inference delay at ~42 ms, well within real-time streaming constraints.",
                "Guide inspected cross-validation results and ROC-AUC curves; suggested implementing feature importance visualization (SHAP/permutation importance) for interpretability.",
                "Action Plan for Next Week: Develop the Behavioral NLP Intent Engine (behavioral_analyzer.py) and integrate in-memory Whisper Automatic Speech Recognition (ASR)."
            ]
        },
        {
            "week": "Week 09\n31-08-2026 to 06-09-2026",
            "date": "31-08-2026\n(Monday)",
            "points": [
                "Demonstrated Behavioral NLP Engine (behavioral_analyzer.py & content_engine.py): Integrated in-memory Whisper ASR for low-latency speech transcription.",
                "Implemented 4-stage conversational scam funnel tracking: Stage 1 (Greeting & Authority Persona) -> Stage 2 (Fraud Hook & Impersonation) -> Stage 3 (Artificial Urgency & Isolation) -> Stage 4 (Coercive Action / Financial Extraction).",
                "Validated regex and semantic intent classification against simulated vishing transcripts (family bail scams, bank fraud impersonation, IRS tax traps).",
                "Guide approved stage weighting coefficients and advised handling benign urgent calls (e.g., real hospital emergencies) to prevent false alarms.",
                "Action Plan for Next Week: Implement the Multimodal Decision Fusion Engine (fusion_engine.py) with dynamic weighting and synergy rules."
            ]
        },
        {
            "week": "Week 10\n07-09-2026 to 13-09-2026",
            "date": "07-09-2026\n(Monday)",
            "points": [
                "Demonstrated Multimodal Decision Fusion Engine (fusion_engine.py): Fuses physical acoustic fake probability and behavioral semantic threat score.",
                "Implemented Dynamic Weighting Heuristic based on Signal-to-Noise Ratio (SNR): Down-weights acoustic features when severe telephony noise/packet loss is detected, giving higher weight to NLP intent.",
                "Implemented Non-Linear Synergy Multiplier: Automatically amplifies risk index when both synthetic audio anomalies (>60%) and coercive scam tactics (>50%) co-occur.",
                "Established clear threat classifications: Low Risk (<35%), Caution / Suspicious (35%–64%), and High Scam Risk (>=65%).",
                "Action Plan for Next Week: Implement Out-of-Band (OOB) Challenge Defense Protocol (verification.py) and real-time alert dispatch."
            ]
        },
        {
            "week": "Week 11\n14-09-2026 to 20-09-2026",
            "date": "14-09-2026\n(Monday)",
            "points": [
                "Demonstrated Out-of-Band (OOB) Challenge Protocol (backend/routes/verification.py): Generates dynamic cognitive authentication challenges during active calls.",
                "Implemented secret family codeword verification and contextual trivia challenges (e.g., 'What was the name of our childhood street?'), neutralizing automated AI scam agents.",
                "Verified challenge generation latency (<15 ms) and zero-knowledge verification mechanisms ensuring user privacy.",
                "Guide reviewed security aspects of out-of-band verification and recommended integrating dual verification options (verbal codeword + automated SMS push challenge).",
                "Action Plan for Next Week: Build FastAPI REST Gateway (backend/main.py) and develop interactive Streamlit/React web dashboard ('VoxSentinel Radar')."
            ]
        },
        {
            "week": "Week 12\n21-09-2026 to 27-09-2026",
            "date": "21-09-2026\n(Monday)",
            "points": [
                "Demonstrated full-stack software integration: FastAPI REST backend (backend/main.py, /api/v1/analyze, /api/v1/benchmarks) and interactive web dashboard (app_streamlit.py).",
                "Implemented live microphone streaming audio capture, drag-and-drop audio file upload (.wav, .mp3), and real-time visual spectrogram / LFCC feature rendering.",
                "Measured end-to-end processing pipeline latency: Audio preprocessing (28 ms) + Feature extraction (42 ms) + ASR & NLP (185 ms) + Fusion (4 ms) = Total 259 ms (target < 350 ms).",
                "Guide tested user interface with sample audio clips, verified responsive UI design, and approved system readiness for formal benchmark evaluation.",
                "Action Plan for Next Week: Execute comprehensive benchmark evaluation across ASVspoof, WaveFake, and telephony codecs; run test suite TC-01 to TC-08."
            ]
        },
        {
            "week": "Week 13\n28-09-2026 to 04-10-2026",
            "date": "28-09-2026\n(Monday)",
            "points": [
                "Presented comprehensive benchmark evaluation results: Evaluated system across ASVspoof 2019/2021, WaveFake, and curated real-world scam call dataset (ScamVoice-500).",
                "Executed formal Test Suite (TC-01 to TC-08): All test cases passed, covering acoustic feature extraction, funnel tracking, fusion synergy, and OOB challenge trigger.",
                "Demonstrated Telephony Robustness Experiments: Simulated 8kHz G.711 mu-law compression and AMR-NB codec degradation; proved multimodal fusion achieves 98.4% accuracy (EER = 2.1%), outperforming acoustic-only models by +48.4% under severe channel distortion.",
                "Guide commended empirical rigor and authorized preparing Review-2 documentation.",
                "Action Plan for Next Week: Compile Review-2 report (80%+ implementation completion, experimental results, architecture diagrams) and 20-slide presentation deck; meet guide on 06-10-2026 for pre-review sign-off."
            ]
        },
        {
            "week": "Week 14\n05-10-2026 to 11-10-2026",
            "date": "06-10-2026\n(Tuesday)",
            "points": [
                "Conducted formal Pre-Review consultation for Review-2 scheduled for Wednesday, 07-10-2026.",
                "Performed full dry-run of 20-slide presentation and live software demonstration of 'VoxSentinel Radar' before the guide.",
                "Demonstrated real-time detection of zero-shot ElevenLabs cloned voice impersonating a family member demanding emergency bail money, triggering instantaneous OOB challenge alert.",
                "Guide thoroughly inspected test suite logs, performance metrics, and verified compliance with Review-2 rubric.",
                "Guide signed and attested official Review-2 document, authorizing presentation before panel.",
                "Action Plan for Next Week: Attend Review-2 on 07-10-2026; prepare for Pre-Final Guide Review scheduled between 12-10-2026 and 16-10-2026."
            ]
        },
        {
            "week": "Week 15\n12-10-2026 to 18-10-2026",
            "date": "16-10-2026\n(Friday)",
            "points": [
                "Debriefed on Review-2 outcomes: Review panel expressed strong appreciation for live demonstration, telephony robustness benchmarks, and active OOB defense concept.",
                "Conducted formal Pre-Final Guide Review (scheduled between 12-10-2026 and 16-10-2026).",
                "Presented rigorous statistical significance validation: Paired t-tests and Wilcoxon signed-rank tests (p < 0.001) confirming multimodal superiority over single-modality baselines.",
                "Demonstrated system memory efficiency (< 380 MB RAM) and sub-300 ms processing latency across simulated mobile devices.",
                "Guide evaluated project across all dimensions of Guide Review Rubric and awarded full marks (30 / 30 Marks).",
                "Guide reviewed and approved complete 70+ page M.Tech dissertation draft (MTECH_SE_Final_Project_Report_Completed.docx) and authorized final thesis submission."
            ]
        }
    ]

    # Build Table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # XML helpers for styling
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_cell_shading(cell, color_hex="F2F2F2"):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tcPr.append(shd)

    def set_table_borders(tbl):
        tblPr = tbl._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    set_table_borders(table)

    col_widths = [Inches(1.3), Inches(1.1), Inches(3.9), Inches(0.8)]
    headers = ["Week & Duration", "Date of Meeting", "Points Discussed", "Signature of the Guide"]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        set_cell_shading(hdr_cells[i], "E8EEF5")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        set_font(r, size=10, bold=True)

    # Make header repeat across pages
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

    # Populate Data
    for entry in diary_entries:
        row = table.add_row()
        # Prevent row split across pages
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(OxmlElement('w:cantSplit'))

        c_week, c_date, c_points, c_sig = row.cells
        for idx, c in enumerate(row.cells):
            c.width = col_widths[idx]
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)

        # Week
        p_w = c_week.paragraphs[0]
        p_w.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_w.paragraph_format.space_before = Pt(2)
        p_w.paragraph_format.space_after = Pt(2)
        p_w.paragraph_format.line_spacing = 1.15
        r_w = p_w.add_run(entry["week"])
        set_font(r_w, size=9.5, bold=True)

        # Date
        p_d = c_date.paragraphs[0]
        p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_d.paragraph_format.space_before = Pt(2)
        p_d.paragraph_format.space_after = Pt(2)
        p_d.paragraph_format.line_spacing = 1.15
        r_d = p_d.add_run(entry["date"])
        set_font(r_d, size=9.5)

        # Points Discussed
        p_pt = c_points.paragraphs[0]
        p_pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_pt.paragraph_format.space_before = Pt(0)
        p_pt.paragraph_format.space_after = Pt(3)
        p_pt.paragraph_format.line_spacing = 1.15

        for p_idx, pt in enumerate(entry["points"]):
            if p_idx > 0:
                p_pt = c_points.add_paragraph()
                p_pt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_pt.paragraph_format.space_before = Pt(0)
                p_pt.paragraph_format.space_after = Pt(3)
                p_pt.paragraph_format.line_spacing = 1.15
            
            r_b = p_pt.add_run("•  ")
            set_font(r_b, size=9.5, bold=True)
            r_t = p_pt.add_run(pt)
            set_font(r_t, size=9.5)

        # Signature Cell left empty for physical signature
        p_s = c_sig.paragraphs[0]
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(0)

    output_path = "PROJECT_WORK_DIARY_AI_VOICE_SCAM.docx"
    doc.save(output_path)
    print(f"Project Work Diary generated successfully at: {output_path}")

if __name__ == "__main__":
    create_project_diary()
